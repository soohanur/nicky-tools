"""
Tabular I/O that preserves the caller's original file.

Why this exists
---------------
The original pipeline reads every input with `CSVReader` and writes a brand new
CSV built from a dict of column -> value. That has three consequences we need to
avoid for boss-ready output:

  1. `.xlsx` uploads are parsed as text and produce garbage (ExcelReader is
     disabled in scraply/src/modules/__init__.py).
  2. Duplicate header names collapse. Real kadaster exports contain `adres`
     twice and `adressen_verblijfsobjectidentificatie` three times, so a
     dict-keyed reader silently drops columns.
  3. All original formatting, column order and sheet structure is lost.

This module reads rows positionally (never by header name), keeps a stable
`row_index` on every row, and produces output by COPYING the original file and
appending new columns to it. Parallel workers may finish rows in any order -
results are written back by `row_index`, never by arrival order.

Some exports put a group-label row above the real header (row 1 "Bron data" /
"Eigenaar", row 2 the column names). Pass `header_row=2` for those; data then
starts on the row after it and results are written back to the same rows.
"""
from __future__ import annotations

import csv
import shutil
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from ..utils.logger import setup_logger

logger = setup_logger('tabular_io')

CSV_ENCODINGS = ['utf-8-sig', 'utf-8', 'cp1252', 'latin-1', 'iso-8859-1']

# Excel caps a cell at 32767 chars; stay well clear.
MAX_CELL_LEN = 3000

# Label written above the new columns when the sheet has a group-label row.
GROUP_LABEL = 'Contactgegevens (company.info)'


def _clean(value: Any) -> str:
    """Normalise a cell to a trimmed string. Ints stay int-looking."""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    # 'NULL' placeholders from earlier tooling must not survive into output.
    return '' if text.upper() == 'NULL' else text


class TabularFile:
    """
    A read/write handle over a CSV or Excel input.

    Usage:
        tf = TabularFile(input_path)                  # header in row 1
        tf = TabularFile(input_path, header_row=2)    # group labels above the header
        rows = tf.read_rows()              # each row: {'row_index': int, headers..., }
        ...
        tf.write_output(out_path, NEW_COLS, {row_index: {col: val}})
    """

    def __init__(self, path: Path, sheet: Optional[str] = None, header_row: int = 1):
        self.path = Path(path)
        if not self.path.exists():
            raise FileNotFoundError(f"File not found: {self.path}")
        self.suffix = self.path.suffix.lower()
        if self.suffix not in ('.csv', '.xlsx', '.xlsm', '.xls'):
            raise ValueError(f"Unsupported file type: {self.suffix}")
        self.sheet = sheet
        # 1-based row holding the column names; data starts on the next row.
        self.header_row = max(1, int(header_row))
        self.headers: List[str] = []
        self._encoding: Optional[str] = None
        self._dialect: Optional[Any] = None
        self._row_count = 0

    # ------------------------------------------------------------------ read

    def read_rows(self) -> List[Dict[str, Any]]:
        if self.suffix == '.csv':
            rows = self._read_csv()
        else:
            rows = self._read_excel()
        self._row_count = len(rows)
        logger.info(f"Read {len(rows)} rows, {len(self.headers)} columns from {self.path.name} "
                    f"(header row {self.header_row})")
        return rows

    def _read_csv(self) -> List[Dict[str, Any]]:
        last_err = None
        for enc in CSV_ENCODINGS:
            try:
                with open(self.path, 'r', encoding=enc, newline='') as f:
                    sample = f.read(64 * 1024)
                    f.seek(0)
                    try:
                        self._dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
                    except Exception:
                        self._dialect = csv.excel
                    reader = csv.reader(f, self._dialect)
                    raw = list(reader)
                self._encoding = enc
                break
            except UnicodeDecodeError as e:
                last_err = e
                continue
        else:
            raise ValueError(f"Could not decode {self.path}: {last_err}")

        hr = self.header_row
        if len(raw) < hr:
            self.headers = []
            return []

        self.headers = [str(h).strip() for h in raw[hr - 1]]
        rows = []
        for i, values in enumerate(raw[hr:]):
            rows.append(self._build_row(i, values))
        return rows

    def _read_excel(self) -> List[Dict[str, Any]]:
        import openpyxl
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        ws = wb[self.sheet] if self.sheet else wb[wb.sheetnames[0]]
        self.sheet = ws.title
        grid = list(ws.iter_rows(values_only=True))
        wb.close()
        hr = self.header_row
        if len(grid) < hr:
            self.headers = []
            return []
        self.headers = [_clean(h) for h in grid[hr - 1]]
        rows = []
        for i, values in enumerate(grid[hr:]):
            rows.append(self._build_row(i, list(values)))
        return rows

    def _build_row(self, index: int, values: Sequence[Any]) -> Dict[str, Any]:
        """
        Positional row. Duplicate header names are preserved by suffixing the
        later occurrences, so no column is ever silently dropped.
        """
        row: Dict[str, Any] = {'row_index': index}
        seen: Dict[str, int] = {}
        cells = []
        for col_i, header in enumerate(self.headers):
            val = _clean(values[col_i]) if col_i < len(values) else ''
            cells.append(val)
            key = header or f"column_{col_i + 1}"
            if key in seen:
                seen[key] += 1
                key = f"{key}__{seen[key]}"
            else:
                seen[key] = 0
            row[key] = val
        row['_cells'] = cells
        return row

    @staticmethod
    def is_blank(row: Dict[str, Any]) -> bool:
        return not any(c for c in row.get('_cells', []))

    # ----------------------------------------------------------------- write

    def write_output(
        self,
        out_path: Path,
        new_columns: List[str],
        results: Dict[int, Dict[str, Any]],
    ) -> int:
        """
        Copy the original file and append `new_columns`, filling each row from
        `results[row_index]`. Rows with no result are left blank (never 'NULL').

        Returns the number of rows populated.
        """
        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        if self.suffix == '.csv':
            # Deliverables go out as Excel even when the upload was a CSV: a CSV
            # round-trip silently rounds 16-digit BAG ids to 15 significant
            # digits and drops the leading zero off every 06... phone number.
            if out_path.suffix.lower() in ('.xlsx', '.xlsm'):
                return self._write_excel_from_csv(out_path, new_columns, results)
            return self._write_csv(out_path, new_columns, results)
        return self._write_excel(out_path, new_columns, results)

    def _write_csv(self, out_path: Path, new_columns, results) -> int:
        with open(self.path, 'r', encoding=self._encoding or 'utf-8-sig', newline='') as f:
            raw = list(csv.reader(f, self._dialect or csv.excel))
        if not raw:
            shutil.copyfile(self.path, out_path)
            return 0

        hr = self.header_row
        out = [list(r) for r in raw[:hr - 1]] + [raw[hr - 1] + list(new_columns)]
        filled = 0
        for i, values in enumerate(raw[hr:]):
            res = results.get(i)
            extra = [_clean(res.get(c)) if res else '' for c in new_columns]
            if res:
                filled += 1
            out.append(list(values) + extra)

        with open(out_path, 'w', encoding='utf-8-sig', newline='') as f:
            csv.writer(f, lineterminator='\r\n').writerows(out)
        logger.info(f"Wrote CSV {out_path.name}: {filled} rows populated")
        return filled

    @staticmethod
    def _looks_numeric(text: str) -> bool:
        """
        True only for values Excel can hold as a number without changing them.

        A leading zero carries meaning here (phone numbers, postcodes) and more
        than 15 digits exceeds Excel's precision, so both stay text.
        """
        if not text or text in ('-', '.', '-.'):
            return False
        body = text[1:] if text[0] in '+-' else text
        if not body or body.count('.') > 1:
            return False
        if not body.replace('.', '', 1).isdigit():
            return False
        digits = body.split('.')[0]
        if len(digits) > 1 and digits[0] == '0':
            return False                       # 0651554613, 07141
        if len(digits.lstrip('0')) > 15:
            return False                       # 1859200000841048
        return True

    def _write_excel_from_csv(self, out_path: Path, new_columns, results) -> int:
        """
        Build a workbook from a CSV upload and append the new columns.

        There is no source workbook to copy here, so the sheet is written from
        scratch: header row bolded and frozen, every value typed deliberately.
        """
        import openpyxl
        from openpyxl.styles import Font

        with open(self.path, 'r', encoding=self._encoding or 'utf-8-sig', newline='') as f:
            raw = list(csv.reader(f, self._dialect or csv.excel))

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = (self.sheet or 'Data')[:31]
        hr = self.header_row

        for r, row in enumerate(raw, start=1):
            values = list(row)
            if r < hr:
                pass                                   # group-label row, as-is
            elif r == hr:
                values = values + list(new_columns)
            else:
                res = results.get(r - hr - 1)
                values = values + [_clean(res.get(c)) if res else '' for c in new_columns]
            for c, value in enumerate(values, start=1):
                text = '' if value is None else str(value).strip()
                if not text:
                    continue
                if len(text) > MAX_CELL_LEN:
                    text = text[:MAX_CELL_LEN]
                cell = ws.cell(row=r, column=c)
                if r >= hr and self._looks_numeric(text):
                    cell.value = float(text) if '.' in text else int(text)
                else:
                    cell.value = text
                    if r > hr:
                        cell.number_format = '@'
                if r == hr:
                    cell.font = Font(bold=True)

        ws.freeze_panes = ws.cell(row=hr + 1, column=1)
        start = len(self.headers) + 1
        for j, name in enumerate(new_columns):
            letter = openpyxl.utils.get_column_letter(start + j)
            ws.column_dimensions[letter].width = 60 if name.upper() == 'NOTE' else 18
        wb.save(out_path)

        filled = sum(1 for i in range(self._row_count) if results.get(i))
        logger.info(f"Wrote Excel (from CSV upload) {out_path.name}: {filled} rows populated")
        return filled

    def _write_excel(self, out_path: Path, new_columns, results) -> int:
        """
        Copy the workbook so ALL original sheets, styles, column widths and
        number formats survive, then append the new columns to the data sheet.
        """
        import openpyxl
        from copy import copy
        from openpyxl.styles import Font

        shutil.copyfile(self.path, out_path)
        wb = openpyxl.load_workbook(out_path)
        ws = wb[self.sheet] if self.sheet in wb.sheetnames else wb[wb.sheetnames[0]]

        start = len(self.headers) + 1
        hr = self.header_row
        last = max(1, len(self.headers))

        # Group-label rows above the header: carry their styling over the new
        # columns and label the block, so it reads as one more section of the
        # same sheet (next to e.g. "Bron data" / "Eigenaar").
        for r in range(1, hr):
            src = ws.cell(row=r, column=last)
            for j in range(len(new_columns)):
                c = ws.cell(row=r, column=start + j)
                if src.has_style:
                    c._style = copy(src._style)
            if r == 1:
                label_src = next((ws.cell(row=1, column=k) for k in range(last, 0, -1)
                                  if ws.cell(row=1, column=k).value not in (None, '')), None)
                lbl = ws.cell(row=1, column=start, value=GROUP_LABEL)
                if label_src is not None and label_src.has_style:
                    lbl._style = copy(label_src._style)

        # Copy the sheet's OWN header style (fill + font + borders) onto the new
        # column headers so they look native — same yellow/bold/etc. as the rest
        # of the header row, instead of a plain pasted-on look.
        hdr_src = ws.cell(row=hr, column=last)
        for j, name in enumerate(new_columns):
            c = ws.cell(row=hr, column=start + j, value=name)
            try:
                if hdr_src.has_style:
                    c._style = copy(hdr_src._style)
                else:
                    c.font = Font(bold=True)
            except Exception:
                c.font = Font(bold=True)

        filled = 0
        for i in range(self._row_count):
            res = results.get(i)
            if not res:
                continue
            excel_row = i + hr + 1                 # rows up to the header, then 1-based
            for j, name in enumerate(new_columns):
                text = _clean(res.get(name))
                if len(text) > MAX_CELL_LEN:
                    text = text[:MAX_CELL_LEN]
                cell = ws.cell(row=excel_row, column=start + j)
                cell.value = text or None          # blank, never the string NULL
                if name.upper().startswith('NUMBER') and text:
                    cell.number_format = '@'       # keep leading zero on 06...
            filled += 1

        wb.save(out_path)
        logger.info(f"Wrote Excel {out_path.name}: {filled} rows populated, "
                    f"columns {start}..{start + len(new_columns) - 1}")
        return filled
